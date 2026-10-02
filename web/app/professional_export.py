# ABOUTME: Renders saved final Markdown as readable Professional DOCX and searchable PDF files.
# ABOUTME: Rejects unsupported content and validates temporary artifacts before publication.

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from html import escape
from pathlib import Path
from typing import cast

from docx import Document
from docx.shared import Inches, Pt
from markdown_it import MarkdownIt
from markdown_it.token import Token
from pypdf import PdfReader
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate

FONT_DIR = Path(__file__).parent / "assets" / "fonts"
MAX_SOURCE_BYTES = 2 * 1024 * 1024
MAX_ARTIFACT_BYTES = 20 * 1024 * 1024


class ExportError(ValueError):
    """A requested format cannot preserve the final document."""


@dataclass(frozen=True)
class Span:
    text: str
    bold: bool = False
    italic: bool = False
    link: bool = False


@dataclass(frozen=True)
class Block:
    kind: str
    level: int
    spans: tuple[Span, ...]
    marker: str = ""

    @property
    def text(self) -> str:
        return "".join(span.text for span in self.spans)


def _inline_spans(children: list[Token]) -> tuple[Span, ...]:
    spans: list[Span] = []
    bold = italic = link = 0
    for child in children:
        kind = child.type
        if kind == "strong_open":
            bold += 1
        elif kind == "strong_close":
            bold -= 1
        elif kind == "em_open":
            italic += 1
        elif kind == "em_close":
            italic -= 1
        elif kind == "link_open":
            link += 1
        elif kind == "link_close":
            link -= 1
        elif kind in {"softbreak", "hardbreak"}:
            spans.append(Span(" " if kind == "softbreak" else "\n", bold > 0, italic > 0, link > 0))
        elif kind == "text":
            if child.content:
                spans.append(Span(child.content, bold > 0, italic > 0, link > 0))
        else:
            raise ExportError(f"Unsupported Markdown: {kind.replace('_', ' ')}")
    return tuple(spans)


def parse_final(markdown: str) -> tuple[Block, ...]:
    if not markdown.strip():
        raise ExportError("Final document is empty")
    if len(markdown.encode("utf-8")) > MAX_SOURCE_BYTES:
        raise ExportError("Final document exceeds the 2 MiB export limit")

    blocks: list[Block] = []
    lists: list[dict[str, int | str | bool]] = []
    current: tuple[str, int, str] | None = None
    for token in MarkdownIt("commonmark").enable("table").parse(markdown):
        kind = token.type
        if kind == "heading_open":
            current = ("heading", int(token.tag[1:]), "")
        elif kind == "paragraph_open":
            if lists and lists[-1]["first"]:
                entry = lists[-1]
                marker = "•" if entry["kind"] == "bullet" else f"{entry['number']}."
                current = (str(entry["kind"]), len(lists), marker)
                entry["first"] = False
            else:
                current = ("paragraph", len(lists), "")
        elif kind == "inline":
            spans = _inline_spans(token.children or [])
            block_kind, level, marker = cast(tuple[str, int, str], current)
            blocks.append(Block(block_kind, level, spans, marker))
        elif kind in {"heading_close", "paragraph_close"}:
            current = None
        elif kind in {"bullet_list_open", "ordered_list_open"}:
            lists.append({"kind": "bullet" if kind == "bullet_list_open" else "ordered",
                          "number": int(token.attrGet("start") or 1) - 1, "first": False})
        elif kind == "list_item_open":
            lists[-1]["number"] = int(lists[-1]["number"]) + 1
            lists[-1]["first"] = True
        elif kind in {"list_item_close", "bullet_list_close", "ordered_list_close"}:
            if kind != "list_item_close":
                lists.pop()
        else:
            raise ExportError(f"Unsupported Markdown: {kind.replace('_', ' ')}")
    return tuple(blocks)


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _validate_content(actual: str, blocks: tuple[Block, ...]) -> None:
    content = _normalize(actual)
    offset = 0
    for block in blocks:
        expected = _normalize(block.text)
        position = content.find(expected, offset)
        if position < 0:
            raise ExportError("Rendered document lost or reordered source text")
        offset = position + len(expected)


def _render_docx(blocks: tuple[Block, ...], path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    for block in blocks:
        if block.kind == "heading":
            style = "Title" if block.level == 1 else f"Heading {block.level - 1}"
        elif block.kind == "bullet":
            style = "List Bullet"
        else:
            style = "Normal"
        paragraph = document.add_paragraph(style=style)
        paragraph.paragraph_format.keep_with_next = block.kind == "heading"
        if block.level and block.kind != "heading":
            paragraph.paragraph_format.left_indent = Inches(0.23 * block.level)
        if block.kind == "ordered":
            paragraph.add_run(f"{block.marker} ")
        for span in block.spans:
            run = paragraph.add_run(span.text)
            run.bold = span.bold or block.kind == "heading"
            run.italic = span.italic
            run.underline = span.link
            run.font.name = "Noto Sans JP" if any(ord(char) >= 0x3000 for char in span.text) else "Noto Sans"
            run.font.size = Pt(16 if block.kind == "heading" and block.level == 1 else 10)
    document.save(str(path))
    _validate_content("\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs), blocks)


def _register_pdf_fonts() -> None:
    for name, filename in (
        ("NotoSans", "NotoSans-Regular.ttf"),
        ("NotoSansBold", "NotoSans-Bold.ttf"),
        ("NotoSansItalic", "NotoSans-Italic.ttf"),
        ("NotoSansBoldItalic", "NotoSans-BoldItalic.ttf"),
        ("NotoJP", "NotoSansJP-VF.ttf"),
    ):
        if name not in pdfmetrics.getRegisteredFontNames():
            font_path = FONT_DIR / filename
            if not font_path.is_file():
                raise ExportError(f"Bundled PDF font is missing: {filename}")
            pdfmetrics.registerFont(TTFont(name, str(font_path)))


def _pdf_font(char: str, bold: bool, italic: bool) -> str:
    if unicodedata.bidirectional(char) in {"R", "AL", "AN"}:
        raise ExportError("PDF cannot safely render right-to-left text; use DOCX or Markdown")
    if unicodedata.category(char).startswith("C"):
        raise ExportError("PDF cannot safely render control or direction characters")
    primary = "NotoSansBoldItalic" if bold and italic else (
        "NotoSansBold" if bold else "NotoSansItalic" if italic else "NotoSans"
    )
    if ord(char) in pdfmetrics.getFont(primary).face.charToGlyph:
        return primary
    if ord(char) in pdfmetrics.getFont("NotoJP").face.charToGlyph:
        return "NotoJP"
    raise ExportError(f"PDF font lacks U+{ord(char):04X}; use DOCX or Markdown")


def _pdf_markup(block: Block) -> str:
    chunks: list[str] = []
    for span in block.spans:
        if span.text == "\n":
            chunks.append("<br/>")
            continue
        segments: list[tuple[str, str]] = []
        for char in span.text:
            font = _pdf_font(char, span.bold or block.kind == "heading", span.italic)
            if segments and segments[-1][0] == font:
                segments[-1] = (font, segments[-1][1] + char)
            else:
                segments.append((font, char))
        for font, text in segments:
            markup = f'<font name="{font}">{escape(text)}</font>'
            chunks.append(f"<u>{markup}</u>" if span.link or span.italic and font == "NotoJP" else markup)
    return "".join(chunks)


def _render_pdf(blocks: tuple[Block, ...], path: Path) -> None:
    _register_pdf_fonts()
    styles = {
        "title": ParagraphStyle("title", fontName="NotoSansBold", fontSize=16, leading=22, spaceAfter=9, keepWithNext=True),
        "heading": ParagraphStyle("heading", fontName="NotoSansBold", fontSize=11, leading=16, spaceBefore=10, spaceAfter=4, keepWithNext=True),
        "body": ParagraphStyle("body", fontName="NotoSans", fontSize=9.5, leading=14, spaceAfter=5),
        "bullet": ParagraphStyle("bullet", fontName="NotoSans", fontSize=9.5, leading=14, spaceAfter=3),
    }
    story: list[Paragraph] = []
    for block in blocks:
        name = "title" if block.kind == "heading" and block.level == 1 else (
            "heading" if block.kind == "heading" else "bullet" if block.kind in {"bullet", "ordered"} else "body"
        )
        style = styles[name]
        if block.kind in {"bullet", "ordered"} or block.level and block.kind == "paragraph":
            style = ParagraphStyle(f"{name}-{block.level}", parent=style, leftIndent=13 * block.level)
        marker = f'<font name="NotoSans">{escape(block.marker)} </font>' if block.marker else ""
        story.append(Paragraph(marker + _pdf_markup(block), style))
    SimpleDocTemplate(str(path), pagesize=letter, leftMargin=54, rightMargin=54,
                      topMargin=50, bottomMargin=50).build(story)
    reader = PdfReader(path)
    _validate_content("\n".join(page.extract_text() or "" for page in reader.pages), blocks)


def export_document(markdown: str, format_name: str, path: Path) -> None:
    """Render and validate one temporary artifact; the caller publishes it atomically."""
    blocks = parse_final(markdown)
    if format_name == "docx":
        _render_docx(blocks, path)
    elif format_name == "pdf":
        _render_pdf(blocks, path)
    else:
        raise ExportError(f"Unsupported export format: {format_name}")
    size = path.stat().st_size
    if not 0 < size <= MAX_ARTIFACT_BYTES:
        raise ExportError("Rendered document is empty or exceeds 20 MiB")
