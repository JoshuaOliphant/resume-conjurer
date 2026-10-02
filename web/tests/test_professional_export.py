# ABOUTME: Verifies Professional exports preserve readable source text and structure.
# ABOUTME: Exercises mixed-script, long-document, and explicit failure behavior with real files.

from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from pypdf import PdfReader

from app import professional_export as professional

SAMPLE = (
    "# José 東京 Αθήνα\n\nSeattle · jose@example.com\n\n## Experience\n\n"
    "**Senior Engineer** — 2020–2026\n\n"
    "- Led €42M program with *focused* work and [portfolio](https://example.com).\n"
    "  - Nested result\n- Next item\n\n1. First ordered\n2. Second ordered\n"
)


def test_mixed_script_exports_preserve_order_and_emphasis(tmp_path: Path) -> None:
    pdf = tmp_path / "resume.pdf"
    docx = tmp_path / "resume.docx"
    professional.export_document(SAMPLE, "pdf", pdf)
    professional.export_document(SAMPLE, "docx", docx)

    pdf_text = "\n".join(page.extract_text() for page in PdfReader(pdf).pages)
    paragraphs = Document(str(docx)).paragraphs
    docx_text = "\n".join(paragraph.text for paragraph in paragraphs)
    for phrase in ("José 東京 Αθήνα", "jose@example.com", "Senior Engineer", "2020–2026", "€42M", "portfolio", "Nested result", "Second ordered"):
        assert phrase in pdf_text
        assert phrase in docx_text
    assert paragraphs[0].style is not None and paragraphs[0].style.name == "Title"
    assert paragraphs[4].style is not None and paragraphs[4].style.name == "List Bullet"
    assert any(run.bold for run in paragraphs[3].runs)
    assert any(run.italic for run in paragraphs[4].runs)
    assert any(run.underline for run in paragraphs[4].runs)
    assert "1. First ordered" in docx_text


def test_long_pdf_flows_across_pages_with_headings_attached(tmp_path: Path) -> None:
    markdown = "# Resume\n\n" + "\n\n".join(
        f"## Role {index}\n\n- Delivered measurable outcome {index} with reliable service."
        for index in range(80)
    )
    pdf = tmp_path / "long.pdf"
    professional.export_document(markdown, "pdf", pdf)
    pages = PdfReader(pdf).pages
    assert len(pages) >= 3
    for page in pages:
        lines = [line.strip() for line in page.extract_text().splitlines() if line.strip()]
        assert not lines[-1].startswith("Role ")
    text = "\n".join(page.extract_text() for page in pages)
    assert text.index("Role 0") < text.index("Role 40") < text.index("Role 79")


def test_wrapped_paragraph_preserves_every_word_in_both_formats(tmp_path: Path) -> None:
    sentence = "Coordinated the migration across twelve teams while preserving service availability and audit evidence."
    markdown = "# Experience\n\n" + " ".join([sentence] * 8)
    pdf = tmp_path / "wrapped.pdf"
    docx = tmp_path / "wrapped.docx"
    professional.export_document(markdown, "pdf", pdf)
    professional.export_document(markdown, "docx", docx)
    pdf_text = " ".join(page.extract_text() for page in PdfReader(pdf).pages).split()
    docx_text = " ".join(paragraph.text for paragraph in Document(str(docx)).paragraphs).split()
    expected = markdown.split()[2:]
    assert pdf_text[1:] == expected
    assert docx_text[1:] == expected


def test_nested_emphasis_keeps_outer_bold_after_inner_close(tmp_path: Path) -> None:
    output = tmp_path / "nested.docx"
    professional.export_document("# Resume\n\n**outer __inner__ outer**", "docx", output)
    runs = Document(str(output)).paragraphs[1].runs
    assert [(run.text, run.bold) for run in runs if run.text] == [
        ("outer ", True), ("inner", True), (" outer", True),
    ]


@pytest.mark.parametrize("markdown, error", [
    ("", "empty"),
    (" " * (professional.MAX_SOURCE_BYTES + 1) + "x", "2 MiB"),
    ("---\n", "Unsupported Markdown"),
    ("# Only heading\n\n<img src='x'>", "Unsupported Markdown"),
    ("# Arabic\n\nمرحبا", "right-to-left"),
    ("# Emoji 😀", "font lacks"),
    ("# Direction \u200e", "direction characters"),
])
def test_pdf_reports_content_it_cannot_preserve(tmp_path: Path, markdown: str, error: str) -> None:
    with pytest.raises(professional.ExportError, match=error):
        professional.export_document(markdown, "pdf", tmp_path / "fail.pdf")


def test_docx_preserves_rtl_text_even_when_pdf_rejects_it(tmp_path: Path) -> None:
    path = tmp_path / "rtl.docx"
    professional.export_document("# Arabic\n\nمرحبا", "docx", path)
    assert "مرحبا" in "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)


def test_missing_bundled_pdf_font_is_explicit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(professional, "FONT_DIR", tmp_path)
    monkeypatch.setattr(professional.pdfmetrics, "getRegisteredFontNames", lambda: [])
    with pytest.raises(professional.ExportError, match="Bundled PDF font is missing"):
        professional.export_document("# Resume", "pdf", tmp_path / "resume.pdf")


def test_unsupported_inline_and_block_markdown_are_explicit() -> None:
    for markdown in (
        "Text ![image](x)", "```python\nprint(1)\n```", "A <b>tag</b>",
        "A | B\n--- | ---\n1 | 2",
    ):
        with pytest.raises(professional.ExportError, match="Unsupported Markdown"):
            professional.parse_final(markdown)


def test_hardbreak_and_escaped_text_survive_pdf(tmp_path: Path) -> None:
    path = tmp_path / "break.pdf"
    professional.export_document("# Resume\n\nOne & two  \nThree < four", "pdf", path)
    text = PdfReader(path).pages[0].extract_text()
    assert "One & two" in text
    assert "Three < four" in text


def test_renderer_rejects_unsupported_format(tmp_path: Path) -> None:
    with pytest.raises(professional.ExportError, match="Unsupported export format"):
        professional.export_document("# Resume", "html", tmp_path / "resume.html")


def test_corrupted_docx_readback_is_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    real_document = professional.Document

    def corrupted_readback(path: str | None = None):
        if path is None:
            return real_document()
        return SimpleNamespace(paragraphs=[SimpleNamespace(text="wrong text")])

    monkeypatch.setattr(professional, "Document", corrupted_readback)
    with pytest.raises(professional.ExportError, match="lost or reordered"):
        professional.export_document("# Resume", "docx", tmp_path / "resume.docx")


def test_text_validation_rejects_missing_or_reordered_content() -> None:
    blocks = professional.parse_final("# Title\n\nFirst\n\nSecond")
    with pytest.raises(professional.ExportError, match="lost or reordered"):
        professional._validate_content("Title Second First", blocks)


@pytest.mark.parametrize("reported_size", [0, professional.MAX_ARTIFACT_BYTES + 1])
def test_invalid_artifact_size_is_rejected(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, reported_size: int
) -> None:
    output = tmp_path / "resume.docx"
    real_stat = Path.stat

    def invalid_stat(path: Path, *args, **kwargs):
        if path == output:
            return SimpleNamespace(st_size=reported_size)
        return real_stat(path, *args, **kwargs)

    monkeypatch.setattr(Path, "stat", invalid_stat)
    with pytest.raises(professional.ExportError, match="empty or exceeds"):
        professional.export_document("# Resume", "docx", output)
